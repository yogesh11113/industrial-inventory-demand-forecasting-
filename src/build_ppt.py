import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def build_review_2_ppt():
    prs = Presentation()
    
    # Define Palette
    NAVY = RGBColor(15, 32, 67)
    BLUE = RGBColor(31, 119, 180)
    DARK_GRAY = RGBColor(50, 50, 50)
    WHITE = RGBColor(255, 255, 255)
    
    def apply_header(slide, title_text):
        top_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.8))
        tf = top_box.text_frame
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = NAVY

    # SLIDE 1: Title Slide
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide1.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = NAVY
    
    title_box = slide1.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(8.4), Inches(3.0))
    tf1 = title_box.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "INDUSTRIAL INVENTORY DEMAND FORECASTING"
    p1.font.size = Pt(30)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    
    p_sub = tf1.add_paragraph()
    p_sub.text = "PBL Project Review-II (30% Progress Milestone)"
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = RGBColor(173, 216, 230)
    p_sub.space_before = Pt(15)
    
    p_author = tf1.add_paragraph()
    p_author.text = "Department of Computer Science Engineering\n1st Year B.E. CSE Student Presentation"
    p_author.font.size = Pt(14)
    p_author.font.color.rgb = WHITE
    p_author.space_before = Pt(30)

    # SLIDE 2: Project Progress Overview
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_header(slide2, "1. Project Progress Overview (30% Milestone)")
    
    tb2 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(8.4), Inches(5.5))
    tf2 = tb2.text_frame
    
    p = tf2.paragraphs[0]
    p.text = "Completed 30% Milestones:"
    p.font.bold = True
    p.font.size = Pt(16)
    
    milestones = [
        "✓ Synthetic Dataset Generation (3,650 daily records across 10 products)",
        "✓ Data Preprocessing & Validation (Missing values check, date parsing, shape: 3650x12)",
        "✓ Exploratory Data Analysis (7 comprehensive charts saved in outputs/graphs/)",
        "✓ Feature Engineering (Temporal features, Lag_1, Lag_7, Rolling_Mean_7, Rolling_Mean_30)",
        "✓ Machine Learning Baseline (Linear Regression with Chronological 80/20 Train/Test Split)",
        "✓ Model Evaluation (Real Metrics: MAE: 22.19, RMSE: 33.40, R² Score: 0.9327)",
        "✓ Basic Reorder Recommendation Logic & Streamlit Web Prototype App"
    ]
    for m in milestones:
        p_m = tf2.add_paragraph()
        p_m.text = m
        p_m.font.size = Pt(13)
        p_m.space_before = Pt(4)
        
    p_rem = tf2.add_paragraph()
    p_rem.text = "Remaining Scope for Review-III & Final Deployment:"
    p_rem.font.bold = True
    p_rem.font.size = Pt(16)
    p_rem.space_before = Pt(15)
    
    remaining = [
        "• Advanced Time-Series Modeling (FB Prophet & LSTM Deep Learning)",
        "• Model Comparison & Hyperparameter Tuning",
        "• Advanced Reorder Point (ROP) & Safety Stock Calculations",
        "• Multi-page Production Dashboard Deployment"
    ]
    for r in remaining:
        p_r = tf2.add_paragraph()
        p_r.text = r
        p_r.font.size = Pt(13)
        p_r.space_before = Pt(3)

    # SLIDE 3: Data Collection & Analysis
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_header(slide3, "2. Data Collection & Analysis")
    
    tb3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(8.4), Inches(5.5))
    tf3 = tb3.text_frame
    
    bullets3 = [
        ("Dataset Description", "Realistic synthetic dataset created for academic reproducibility."),
        ("Sample Size", "3,650 total records spanning 365 days (2025-01-01 to 2025-12-31) for 10 distinct products."),
        ("Core Features", "Date, Product_ID, Category, Sales_Quantity, Consumption, Current_Stock, Lead_Time_Days, Supplier, Unit_Price, Stock_Level, Demand."),
        ("Demand Patterns", "Incorporates baseline demand, annual sinusoidal seasonality, weekday/weekend variations, and realistic random noise."),
        ("Preprocessing Results", "Zero missing values, zero duplicates, sorted chronologically, normalized stock boundaries.")
    ]
    for title, desc in bullets3:
        p_t = tf3.add_paragraph()
        p_t.text = f"• {title}: {desc}"
        p_t.font.size = Pt(14)
        p_t.space_before = Pt(10)

    # SLIDE 4: Methodology & Design
    slide4 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_header(slide4, "3. Methodology & Design")
    
    tb4 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(8.4), Inches(5.5))
    tf4 = tb4.text_frame
    
    bullets4 = [
        ("Data Pipeline", "Raw CSV -> Preprocessing & Cleaning -> Feature Extraction -> Chronological Split -> Model Training"),
        ("Preventing Data Leakage", "All lag (Lag_1, Lag_7) and rolling statistics (Rolling_Mean_7, 30) are strictly computed using shifted past historical demand."),
        ("Time-Based Train/Test Split", "Applied an 80% past chronological training set and 20% future test set (no random shuffling to preserve time ordering)."),
        ("Reorder Decision Logic", "Triggers REORDER REQUIRED if Current_Stock < (Predicted_Demand * Lead_Time * 1.2 Safety Margin).")
    ]
    for title, desc in bullets4:
        p_t = tf4.add_paragraph()
        p_t.text = f"• {title}: {desc}"
        p_t.font.size = Pt(14)
        p_t.space_before = Pt(12)

    # SLIDE 5: System Architecture & Module Design
    slide5 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_header(slide5, "4. System Architecture & Module Design")
    
    tb5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(8.4), Inches(5.5))
    tf5 = tb5.text_frame
    
    modules = [
        ("generate_data.py", "Synthesizes multi-product historical dataset"),
        ("data_preprocessing.py", "Cleans, formats, and validates dataset integrity"),
        ("eda.py", "Generates 7 EDA analytical plots for demand & inventory distribution"),
        ("feature_engineering.py", "Constructs temporal, lag, and rolling features without data leakage"),
        ("model_training.py", "Trains Linear Regression model and exports evaluation metrics & predictions"),
        ("prediction.py", "Loads saved joblib model for real-time demand inference"),
        ("reorder_recommendation.py", "Executes stock replenishment recommendation algorithm"),
        ("app/app.py", "Streamlit web interface for interactive product demand lookup")
    ]
    for mod, desc in modules:
        p_m = tf5.add_paragraph()
        p_m.text = f"• src/{mod} — {desc}"
        p_m.font.size = Pt(13)
        p_m.space_before = Pt(6)

    # SLIDE 6: Implementation Environment
    slide6 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_header(slide6, "5. Implementation Environment")
    
    tb6 = slide6.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(8.4), Inches(5.5))
    tf6 = tb6.text_frame
    
    env_items = [
        ("Programming Language", "Python 3.12 / 3.14"),
        ("Development Environment", "VS Code / Jupyter Notebook"),
        ("Data Processing & Math", "Pandas, NumPy"),
        ("Machine Learning Library", "Scikit-Learn (LinearRegression)"),
        ("Model Persistence", "Joblib"),
        ("Visualization Libraries", "Matplotlib, Seaborn"),
        ("Web Application Framework", "Streamlit")
    ]
    for label, val in env_items:
        p_e = tf6.add_paragraph()
        p_e.text = f"• {label}: {val}"
        p_e.font.size = Pt(14)
        p_e.space_before = Pt(10)

    # SLIDE 7: Implementation Progress & Results
    slide7 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_header(slide7, "6. Implementation Progress & Evaluation Results")
    
    tb7 = slide7.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(8.4), Inches(5.5))
    tf7 = tb7.text_frame
    
    results = [
        ("Actual Model Performance (Linear Regression Baseline)", ""),
        ("Mean Absolute Error (MAE)", "22.19 units"),
        ("Mean Squared Error (MSE)", "1115.77 units²"),
        ("Root Mean Squared Error (RMSE)", "33.40 units"),
        ("R² Score (Coefficient of Determination)", "0.9327 (93.27% Variance Explained)"),
        ("Sample Recommendation Output", "Product: PRD_001 | Predicted Demand: 120 | Current Stock: 80 | Status: REORDER REQUIRED | Recommended Qty: 640 units")
    ]
    for label, val in results:
        p_r = tf7.add_paragraph()
        if val == "":
            p_r.text = f"• {label}"
            p_r.font.bold = True
            p_r.font.size = Pt(15)
        else:
            p_r.text = f"   - {label}: {val}"
            p_r.font.size = Pt(13)
        p_r.space_before = Pt(6)

    # SLIDE 8: Current Challenges & Future Direction
    slide8 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_header(slide8, "7. Current Challenges & Future Direction")
    
    tb8 = slide8.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(8.4), Inches(5.5))
    tf8 = tb8.text_frame
    
    p = tf8.paragraphs[0]
    p.text = "Current Technical & Modeling Challenges:"
    p.font.bold = True
    p.font.size = Pt(15)
    
    challenges = [
        "1. Capturing Complex Non-Linear Seasonality (Linear Regression has baseline limits)",
        "2. Handling Sudden Supply Chain Lead Time Fluctuations",
        "3. Optimizing Safety Stock Parameters dynamically"
    ]
    for c in challenges:
        p_c = tf8.add_paragraph()
        p_c.text = c
        p_c.font.size = Pt(13)
        p_c.space_before = Pt(4)
        
    p_f = tf8.add_paragraph()
    p_f.text = "Future Action Plan for Review-III & Final Review:"
    p_f.font.bold = True
    p_f.font.size = Pt(15)
    p_f.space_before = Pt(12)
    
    future = [
        "• Implementation of Prophet & LSTM Neural Network Models",
        "• Comparative Analysis (Linear Regression vs Prophet vs LSTM)",
        "• Dynamic Safety Stock & Economic Order Quantity (EOQ) Integration",
        "• Deployment of Full Interactive Dashboard"
    ]
    for f_item in future:
        p_fi = tf8.add_paragraph()
        p_fi.text = f_item
        p_fi.font.size = Pt(13)
        p_fi.space_before = Pt(4)

    # SLIDE 9: Thank You
    slide9 = prs.slides.add_slide(prs.slide_layouts[6])
    bg9 = slide9.background
    fill9 = bg9.fill
    fill9.solid()
    fill9.fore_color.rgb = NAVY
    
    ty_box = slide9.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(7.0), Inches(2.5))
    tf9 = ty_box.text_frame
    p9 = tf9.paragraphs[0]
    p9.text = "THANK YOU!"
    p9.alignment = PP_ALIGN.CENTER
    p9.font.size = Pt(36)
    p9.font.bold = True
    p9.font.color.rgb = WHITE
    
    p9_sub = tf9.add_paragraph()
    p9_sub.text = "Questions & Feedback Welcome\nIndustrial Inventory Demand Forecasting — Review-II"
    p9_sub.alignment = PP_ALIGN.CENTER
    p9_sub.font.size = Pt(18)
    p9_sub.font.color.rgb = RGBColor(173, 216, 230)
    p9_sub.space_before = Pt(15)

    os.makedirs("outputs", exist_ok=True)
    ppt_path = "outputs/Review_II_Presentation.pptx"
    prs.save(ppt_path)
    print(f"Presentation saved successfully to '{ppt_path}'")

if __name__ == "__main__":
    build_review_2_ppt()
